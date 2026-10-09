"""Protected trial dashboard. OneDrive data stays outside the public repository."""
import configparser
import hashlib
import io
import os
import re
import secrets
import shutil
import subprocess
import threading
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from functools import wraps

from flask import Flask, abort, jsonify, redirect, render_template_string, request, send_file, session
from openpyxl import load_workbook
from werkzeug.security import check_password_hash, generate_password_hash

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('DATA_DIR', '/app/storage'))
INTERVAL = max(60, int(os.environ.get('SYNC_INTERVAL_SECONDS', '300')))
MAX_BYTES = 32 * 1024 * 1024
CONFIG = DATA / 'rclone.conf'
SEED = Path(os.environ.get('RCLONE_CONFIG_SEED', '/etc/secrets/rclone.conf'))
REMOTE = os.environ.get('ONEDRIVE_REMOTE', 'onedrive')
FILE_PATH = os.environ.get('ONEDRIVE_FILE_PATH', '').strip()
SECRET = os.environ.get('SECRET_KEY', '')
PASSWORD = os.environ.get('DASHBOARD_PASSWORD', '')
PASSWORD_HASH = generate_password_hash(PASSWORD) if len(PASSWORD) >= 16 else None
app = Flask(__name__, static_folder=None)
app.config.update(SECRET_KEY=SECRET or secrets.token_hex(32),
                  SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', 'true').lower() != 'false',
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=8), MAX_CONTENT_LENGTH=8192)
lock = threading.RLock()
sync_lock = threading.Lock()
state = {'connected': False, 'digest': None, 'last_success': None,
         'last_checked': None, 'error': 'not_configured', 'interval_seconds': INTERVAL}
snapshot = None
failures = []
failure_lock = threading.Lock()


def now():
    return datetime.now(timezone.utc).isoformat()


def credentials_ready():
    return PASSWORD_HASH is not None and len(SECRET) >= 32


def protected(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not credentials_ready():
            return 'Set DASHBOARD_PASSWORD (16+ characters) and SECRET_KEY (32+ characters) in Render.', 503
        if not session.get('authenticated'):
            if request.path.startswith('/api/'):
                return jsonify(error='authentication_required'), 401
            return redirect('/login')
        return fn(*args, **kwargs)
    return wrapped


@app.after_request
def headers(response):
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; font-src 'self' data:; img-src 'self' data: blob:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    return response


LOGIN = '''<!doctype html><html lang="th"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PLI Trial Dashboard — Sign in</title>
<style>body{margin:0;background:#0c1c32;color:#eef5fc;font:18px Tahoma,sans-serif;display:grid;min-height:100vh;place-items:center}.card{width:min(390px,85vw);padding:32px;background:#142b46;border-radius:20px}h1{font-size:26px}p{line-height:1.6;color:#bdd2e3}input,button{box-sizing:border-box;width:100%;padding:15px;border-radius:10px;margin-top:12px;font:inherit}button{background:#31d4b1;border:0;color:#092735;cursor:pointer}.error{color:#ffb5ab}</style>
<div class="card"><h1>PLI Trial Dashboard</h1><p>ทะเบียนขอเบิกและติดตามตัวอย่าง<br>Sample requests &amp; trial follow-up</p><form method="post"><input type="hidden" name="csrf" value="{{ csrf }}"><label for="password">รหัสผ่าน / Access password</label><input id="password" type="password" name="password" required autocomplete="current-password" maxlength="512"><button>เข้าสู่ระบบ / Sign in</button></form>{% if error %}<p class="error">{{ error }}</p>{% endif %}</div></html>'''


@app.route('/login', methods=['GET', 'POST'])
def login():
    if not credentials_ready():
        return 'Dashboard access has not been configured in Render.', 503
    session.setdefault('csrf', secrets.token_urlsafe(32))
    error = ''
    if request.method == 'POST':
        supplied = request.form.get('csrf', '')
        if not supplied or not secrets.compare_digest(supplied, session['csrf']):
            abort(400)
        with failure_lock:
            failures[:] = [t for t in failures if time.monotonic() - t < 60]
            # Global rate limit avoids trusting spoofable proxy headers.
            if len(failures) >= 20:
                return 'Please wait one minute before trying again.', 429
            failures.append(time.monotonic())
        if check_password_hash(PASSWORD_HASH, request.form.get('password', '')[:512]):
            session.clear()
            session['authenticated'] = True
            session['csrf'] = secrets.token_urlsafe(32)
            session.permanent = True
            return redirect('/')
        error = 'รหัสผ่านไม่ถูกต้อง / Incorrect password'
    return render_template_string(LOGIN, csrf=session['csrf'], error=error)


@app.post('/logout')
@protected
def logout():
    if not secrets.compare_digest(request.form.get('csrf', ''), session.get('csrf', '')):
        abort(400)
    session.clear()
    return redirect('/login')


@app.get('/healthz')
def health():
    return jsonify(status='ok')


@app.get('/')
@app.get('/th')
@protected
def thai():
    return send_file(ROOT / 'web/th.html')


@app.get('/en')
@protected
def english():
    return send_file(ROOT / 'web/en.html')


@app.get('/live.js')
@protected
def live_script():
    return send_file(ROOT / 'web/live.js')


@app.get('/mobile.css')
@protected
def mobile_styles():
    return send_file(ROOT / 'web/mobile.css', mimetype='text/css')


@app.get('/api/status')
@protected
def sync_status():
    with lock:
        result = dict(state)
    result['csrf'] = session['csrf']
    return jsonify(result)


@app.get('/api/workbook')
@protected
def workbook():
    with lock:
        data, digest = snapshot, state['digest']
    if data is None:
        return jsonify(error='workbook_unavailable'), 503
    response = send_file(io.BytesIO(data), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         download_name='Sample Trial Record.xlsx', as_attachment=False)
    response.headers['ETag'] = '"' + digest + '"'
    return response


def validate_workbook(data):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Invalid workbook size')
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if 'Sample Trial Record' not in book.sheetnames:
            raise ValueError('Required worksheet missing')
    finally:
        book.close()


def prepare_config():
    DATA.mkdir(parents=True, exist_ok=True, mode=0o700)
    if not CONFIG.exists() and SEED.is_file():
        shutil.copyfile(SEED, CONFIG)
        CONFIG.chmod(0o600)
    if not CONFIG.is_file() or not FILE_PATH:
        return False
    parsed = configparser.ConfigParser(interpolation=None)
    parsed.read(CONFIG)
    if REMOTE not in parsed or parsed[REMOTE].get('type') != 'onedrive':
        raise ValueError('An authenticated OneDrive remote is required')
    if not FILE_PATH.lower().endswith('.xlsx') or FILE_PATH.startswith('-') or '\n' in FILE_PATH:
        raise ValueError('Configure a relative XLSX path in OneDrive')
    return True


def safe_download_diagnostic(stderr, returncode):
    """Return only predefined words and machine codes, never raw command text."""
    detail = (stderr or b'').decode('utf-8', errors='replace').lower()
    vocabulary = set('failed error fatal creating create file system config configuration section remote onedrive drive personal business invalid unsupported unknown option flag hash type auto quickxor sha1 token refresh expired expiry authentication authorization access denied forbidden permission scope client grant json parse parsing decode decoding unexpected character end input eof empty missing not found directory object item download downloading copy copying transfer checksum mismatch corrupted size network timeout connection connect refused reset dns lookup tls certificate resolve resolving host throttled rate limit too many requests service unavailable unauthorized encrypted password decrypt decryption malformed'.split())
    # Exact whole-word membership: arbitrary identifiers, URLs, tokens and paths are excluded.
    words = [word for word in re.findall(r'[a-z]+', detail) if word in vocabulary]
    codes = sorted(set(re.findall(r'\baadsts[0-9]{5,8}\b', detail)))
    statuses = sorted(set(re.findall(r'(?:status(?: code)?|http)[ :/]+([45][0-9]{2})\b', detail)))
    return {'exit_code': int(returncode), 'error_words': words[-40:],
            'microsoft_codes': codes[:5], 'http_statuses': statuses[:5]}


def sync_once():
    global snapshot
    if not sync_lock.acquire(blocking=False):
        return
    target = DATA / 'incoming.xlsx'
    failure_code = 'config_invalid'
    try:
        with lock:
            state['last_checked'] = now()
            state['download_diagnostic'] = None
        if not prepare_config():
            with lock:
                state['connected'] = False
                state['error'] = 'not_configured'
            return
        failure_code = 'download_failed'
        # Fixed argument list, no shell. Only downloads the configured source file.
        proc = subprocess.run(['rclone', 'copyto', REMOTE + ':' + FILE_PATH, str(target),
                               '--config', str(CONFIG),
                               '--retries', '2', '--low-level-retries', '2', '--log-level', 'ERROR'],
                              stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=150)
        if proc.returncode:
            diagnostic = safe_download_diagnostic(getattr(proc, 'stderr', b''), proc.returncode)
            with lock:
                state['download_diagnostic'] = diagnostic
            app.logger.warning('OneDrive download diagnostic: %s', diagnostic)
            # Inspect in memory only; expose fixed categories, never raw output.
            detail = (proc.stderr or b'').decode('utf-8', errors='replace').lower()
            if any(term in detail for term in ('invalid_grant', 'unauthorized', 'invalid_client', 'token expired', "couldn't fetch token")):
                failure_code = 'authentication_failed'
            elif any(term in detail for term in ('directory not found', 'object not found', 'itemnotfound')):
                failure_code = 'file_not_found'
            raise ValueError('OneDrive download failed')
        if not target.is_file() or target.stat().st_size > MAX_BYTES:
            failure_code = 'download_empty_or_oversized'
            raise ValueError('OneDrive download failed')
        failure_code = 'workbook_invalid'
        data = target.read_bytes()
        validate_workbook(data)
        digest = hashlib.sha256(data).hexdigest()
        with lock:
            snapshot = data
            state.update(connected=True, digest=digest, last_success=now(), error=None, failure_stage=None)
    except subprocess.TimeoutExpired:
        with lock:
            state.update(connected=False, error='sync_failed', failure_stage='download_timeout')
        app.logger.warning('OneDrive sync: download_timeout')
    except Exception as exc:
        # Never expose command output, tokens, paths, or workbook content in logs/API.
        with lock:
            state['connected'] = False
            state['error'] = 'sync_failed'
            state['failure_stage'] = failure_code
        app.logger.warning('OneDrive sync: %s (%s)', failure_code, type(exc).__name__)
    finally:
        target.unlink(missing_ok=True)
        sync_lock.release()


def sync_loop():
    while True:
        sync_once()
        time.sleep(INTERVAL)


if os.environ.get('DISABLE_SYNC') != 'true':
    threading.Thread(target=sync_loop, daemon=True, name='onedrive-reader').start()
